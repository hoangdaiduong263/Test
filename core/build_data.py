"""Dựng gói dữ liệu D (giống buildData_ trong gas/Code.gs) từ hai file xlsx tải về, không cần Apps Script.
python3 core/build_data.py <VolumeTracking.xlsx> <Linehaul_Data.xlsx> <data.js cũ (lấy REF)> <data.js mới> [ngày cuối, ví dụ 2026-09-18]"""
import sys, json, re, math, datetime as dt_
import openpyxl

VOL_X, LH_X, OLD, OUT = sys.argv[1:5]
END = sys.argv[5] if len(sys.argv) > 5 else '9999'
VEH = ['1T9', '5T', '1T25', 'VAN', '8T', 'KHAC']
EPOCH = datetime = dt_.datetime(1899, 12, 30)

def num(v):
    if v is None or v == '': return 0
    if isinstance(v, dt_.datetime): return round((v - EPOCH).total_seconds() / 86400)
    if isinstance(v, str):
        m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', v)
        if m: return round((dt_.datetime(int(m[1]), int(m[2]), int(m[3])) - EPOCH).total_seconds() / 86400)
        if re.match(r'^\d{1,2}:\d{2}', v): return 0
        try: return float(v)
        except ValueError: return 0
    try: return float(v)
    except (TypeError, ValueError): return 0
def dstr(v): return v.strftime('%Y-%m-%d') if isinstance(v, dt_.datetime) else str(v or '')[:10]
def idx(h, names):
    m = {}
    for i, x in enumerate(h):
        k = str(x).strip() if x is not None else ''
        if k not in m: m[k] = i
    return {k: next((m[n] for n in ns if n in m), -1) for k, ns in names.items()}
def vtype(v):
    v = str(v or '').upper()
    for k in ('1T9', '1T25', '5T', '8T', 'VAN'):
        if k in v: return k
    return 'KHAC'
def median(a):
    a = sorted(x for x in a if x is not None)
    if not a: return None
    m = len(a) >> 1
    return a[m] if len(a) % 2 else (a[m - 1] + a[m]) / 2
def sheet(wb, name):
    norm = lambda s: re.sub(r'[^a-z0-9]', '', str(s).lower())
    for ws in wb.worksheets:
        if ws.title == name or norm(ws.title) == norm(name): return ws
    return None
def rows(ws): return list(ws.iter_rows(values_only=True))

vwb = openpyxl.load_workbook(VOL_X, read_only=True, data_only=True)
lwb = openpyxl.load_workbook(LH_X, read_only=True, data_only=True)

# 1. đơn theo điểm theo ngày
vv = rows(sheet(vwb, 'Transformed'))
vi = idx(vv[0], dict(region=['region'], date=['cdate'], hub=['fm_hub_name', 'fm_hub_cover'], vs=['virtual_station_name'], vol=['shipment_count', 'total_vol'], bk=['bulky_count', 'bulky_order']))
vol, dateSet, hubVol, regionOf = {}, set(), {}, {}
for r in vv[1:]:
    vs = str(r[vi['vs']] or '').strip()
    if not vs: continue
    d = dstr(r[vi['date']])
    if not d: continue
    q, bq = num(r[vi['vol']]), num(r[vi['bk']])
    if q <= 0: continue
    dateSet.add(d); regionOf[vs] = str(r[vi['region']] or '').strip()
    k = vs + '|' + d; vol.setdefault(k, [0, 0]); vol[k][0] += q; vol[k][1] += bq
    hub = str(r[vi['hub']] or '').strip()
    if hub: hubVol[vs + '|' + hub] = hubVol.get(vs + '|' + hub, 0) + q
dates = [d for d in sorted(dateSet) if d <= END][-92:]
dayIx = {d: i for i, d in enumerate(dates)}; ND = len(dates)

# 2. loại ngày
dtv = [0] * ND; warn = []
csh = sheet(vwb, 'Config [Report]')
if csh:
    cv = rows(csh); ci = idx(cv[0], dict(d=['DATE'], t=['TYPE']))
    for r in cv[1:]:
        if ci['d'] < 0 or ci['t'] < 0: break
        dd = dstr(r[ci['d']])
        if dd not in dayIx: continue
        ty = str(r[ci['t']] or '').upper()
        dtv[dayIx[dd]] = 2 if ty.startswith('CP') else (1 if ty.startswith('MINI') else 0)
else: warn.append('Không thấy sheet loại ngày')

# 3. chuyến
lv = rows(sheet(lwb, 'Data'))
li = idx(lv[0], dict(date=['Ngày'], code=['Mã chuyến'], seq=['Thứ tự dừng'], name=['Tên điểm'], type=['Loại điểm'], up=['Đơn lên tại điểm', 'TO lên tại điểm'],
    veh=['Loại xe'], load=['Giờ tải hàng (loading_time)'], src=['Nguồn chuyến'], vendor=['Nhà thầu'], dn=['Đơn xuống tại điểm', 'TO xuống tại điểm'],
    toN=['Số TO (API đơn)'], arr=['Giờ đến điểm'], dep=['Giờ đi điểm'], arrK=['Giờ đến KH'], plate=['Biển số'], km=['Quãng đường (km)']))
G = lambda r, k: r[li[k]] if li[k] >= 0 and li[k] < len(r) else None
def minOf(v, d0):
    if not isinstance(v, dt_.datetime): return None
    dd = (dt_.datetime.strptime(v.strftime('%Y-%m-%d'), '%Y-%m-%d') - dt_.datetime.strptime(d0, '%Y-%m-%d')).days
    return dd * 1440 + v.hour * 60 + v.minute
rowsOf, order = {}, []
for r in lv[1:]:
    c = str(G(r, 'code') or '').strip()
    if not c: continue
    if c not in rowsOf: rowsOf[c] = []; order.append(c)
    rowsOf[c].append(r)
dayOfTrip = {}
for c in order:
    rs = sorted(rowsOf[c], key=lambda r: num(G(r, 'seq')))
    d0 = dstr(G(rs[0], 'date'))
    for r in rs:
        ty = str(G(r, 'type') or '').strip()
        if ty in ('SOC', 'Hub'): continue
        av = G(r, 'arr')
        if isinstance(av, dt_.datetime): d0 = dstr(av)
        break
    dayOfTrip[c] = d0
trips, lhSeen = {}, set()
for c in order:
    for r in rowsOf[c]:
        d2 = dayOfTrip[c]
        if d2 not in dayIx: continue
        lhSeen.add(d2)
        if c not in trips:
            km = G(r, 'km')
            trips[c] = dict(d=d2, veh=vtype(G(r, 'veh')), stops=[], socs={}, upAll=0, upD2S=0, upHub=0, firstType=None,
                adhoc=str(G(r, 'src') or '').strip().lower() == 'adhoc', inh=str(G(r, 'vendor') or '').strip().lower() == 'in-house',
                path=[], plate=str(G(r, 'plate') or '').strip(), km=(round(num(km) * 10) / 10 if km not in (None, '') else None))
        t = trips[c]; typ = str(G(r, 'type') or '').strip(); up = num(G(r, 'up')); name = str(G(r, 'name') or '').strip(); seq = num(G(r, 'seq'))
        if t['firstType'] is None or seq == 1: t['firstType'] = typ
        t['path'].append([seq, name, 2 if typ == 'SOC' else (1 if typ == 'Hub' else 0), up, num(G(r, 'dn')) if li['dn'] >= 0 else 0,
            minOf(G(r, 'arr'), d2), minOf(G(r, 'dep'), d2), minOf(G(r, 'arrK'), d2), num(G(r, 'toN')) if li['toN'] >= 0 else 0])
        t['upAll'] += up
        if typ == 'SOC': t['socs'][name] = 1
        elif typ == 'Hub': t['upHub'] += up
        else: t['upD2S'] += up; t['stops'].append(dict(vs=name, up=up, load=G(r, 'load')))

# 3b. chuyến Hub tự đi lấy hàng (PBT): lên ở điểm D2S, không ghé SOC nào, trả ở Hub → không phải D2S, loại khỏi chuyến D2S.
#     Đơn của điểm hôm đó trừ phần đi Hub (theo tỷ lệ đơn lên trên chuyến thật). Chuyến có ghé SOC mà ghi xuống 0 vẫn là D2S (coi như xuống hết ở SOC đó).
hubUp, socUp, pbt = {}, {}, {}
for code, t in list(trips.items()):
    if not t['stops']: continue
    toHub = not t['socs'] and any(p[2] == 1 for p in t['path'])
    for s in t['stops']:
        k = s['vs'] + '|' + t['d']; (hubUp if toHub else socUp)[k] = (hubUp if toHub else socUp).get(k, 0) + s['up']
    if toHub:
        for s in t['stops']: pbt[s['vs']] = pbt.get(s['vs'], 0) + 1
        del trips[code]
if pbt: warn.append('Loại ' + str(sum(pbt.values())) + ' lượt lấy hàng về Hub (PBT, không phải D2S): ' + ', '.join(f'{k} {n}' for k, n in sorted(pbt.items(), key=lambda x: -x[1])))
def d2sFrac(k):
    h = hubUp.get(k, 0)
    if h <= 0: return 1
    s_ = socUp.get(k, 0); return s_ / (s_ + h)

# 4. gom theo điểm – ngày
per, pairCount = {}, {}
for code, t in trips.items():
    if not t['stops']: continue
    share = t['upD2S'] / t['upAll'] if t['upAll'] > 0 else 1
    nSoc = len(t['socs']); vIx = VEH.index(t['veh']) if t['veh'] in VEH else len(VEH) - 1
    names = [s['vs'] for s in t['stops']]
    for a in range(len(names)):
        for b in range(a + 1, len(names)):
            kk = '||'.join(sorted([names[a], names[b]])); pairCount[kk] = pairCount.get(kk, 0) + 1
    for s in t['stops']:
        frac = share * s['up'] / t['upD2S'] if t['upD2S'] > 0 else share / len(t['stops'])
        k = s['vs'] + '|' + t['d']
        p = per.setdefault(k, dict(tr=[0] * 6, tc=[], to=0, socs={}, multi=[0, 0], loads=[], types={}, ad=0, ih=0))
        p['tr'][vIx] += frac
        if t['adhoc']: p['ad'] += frac
        if t['inh']: p['ih'] += frac
        if code not in p['tc']: p['tc'].append(code)
        p['to'] += s['up']
        for n in t['socs']: p['socs'][n] = 1
        p['multi'][1] += 1
        if nSoc >= 2: p['multi'][0] += 1
        if isinstance(s['load'], dt_.datetime): p['loads'].append(s['load'])
        tt = 'T4' if t['firstType'] == 'SOC' else ('T3' if t['upHub'] > 0 else ('T1' if len(t['stops']) == 1 else 'T2'))
        p['types'][tt] = p['types'].get(tt, 0) + 1

# 5. cụm ghép
parent = {}
def find(x):
    while parent[x] != x: parent[x] = parent[parent[x]]; x = parent[x]
    return x
def union(a, b):
    parent.setdefault(a, a); parent.setdefault(b, b); ra, rb = find(a), find(b)
    if ra != rb: parent[ra] = rb
for k, n in pairCount.items():
    if n >= 3: a, b = k.split('||'); union(a, b)
clusterName, clusterSeq = {}, {}
for v in list(parent):
    root = find(v); reg = (regionOf.get(v) or 'X')[:4]
    if root not in clusterName: clusterSeq[reg] = clusterSeq.get(reg, 0) + 1; clusterName[root] = reg + '-X' + str(clusterSeq[reg])

# 6. danh sách điểm
vsList = sorted({k.split('|')[0] for k in vol})
hubOf = {}
for k, q in hubVol.items():
    vs, hub = k.split('|', 1)
    if vs not in hubOf or q > hubOf[vs][1]: hubOf[vs] = [hub, q]
hubCount = {}
for vs, h in hubOf.items(): key = (regionOf.get(vs) or '') + '|' + h[0]; hubCount[key] = hubCount.get(key, 0) + 1
typeDays = [dtv.count(0), dtv.count(1), dtv.count(2)]
r3 = lambda x: round(x * 1000) / 1000
S = []
for vs in vsList:
    v, b, tr, tc, to, ad, ih, runT = [], [], [], [], [], [], [], [0, 0, 0]
    nsList, mdN, mdD, waves, typeCount, socCount, wvd, lt1 = [], 0, 0, [], {}, {}, [], []
    for i in range(ND):
        q = vol.get(vs + '|' + dates[i], [0, 0]); f = d2sFrac(vs + '|' + dates[i]); q = [q[0] * f, q[1] * f]; v.append(round(q[0])); b.append(round(q[1]))
        if q[0] > 0: runT[dtv[i]] += 1
        p = per.get(vs + '|' + dates[i])
        if p:
            tr.append([r3(x) for x in p['tr']]); tc.append(p['tc']); to.append(round(p['to'])); ad.append(r3(p['ad'])); ih.append(r3(p['ih']))
            nsList.append(len(p['socs']))
            for n in p['socs']: socCount[n] = socCount.get(n, 0) + 1
            mdN += p['multi'][0]; mdD += p['multi'][1]
            ls = sorted(p['loads']); w = 1 if ls else 0
            for j in range(1, len(ls)):
                if (ls[j] - ls[j - 1]).total_seconds() > 2 * 3600: w += 1
            if w: waves.append(w)
            wvd.append(w or None); lt1.append(ls[0].strftime('%H:%M') if ls else None)
            for t_, n in p['types'].items(): typeCount[t_] = typeCount.get(t_, 0) + n
        else:
            tr.append(0 if q[0] > 0 else None); tc.append(None); to.append(0); ad.append(None); ih.append(None); wvd.append(None); lt1.append(None)
    if not any(x > 0 for x in v): continue
    tt = sorted(typeCount, key=lambda x: -typeCount[x])[0] if typeCount else ''
    hub = hubOf[vs][0] if vs in hubOf else ''
    socs = sorted(socCount, key=lambda n: -socCount[n])[:8]
    S.append(dict(n=vs, R=regionOf.get(vs, ''), k='SPC' if 'SPC' in vs.upper() else 'Seller', h=hub, v=v, b=b, tr=tr, tc=tc, to=to, ad=ad, ih=ih,
        on=[1 if typeDays[t_] > 0 and runT[t_] >= 0.5 * typeDays[t_] else 0 for t_ in range(3)],
        ns=median(nsList) or 1, nsMed=median(nsList), nsMax=max(nsList) if nsList else None, mdSh=round(mdN / mdD * 100) / 100 if mdD else 0,
        wv=round(median(waves)) if waves else 1, wvd=wvd, lt1=lt1, soc=socs, socN=[socCount[n] for n in socs], tpd=None, tt=tt,
        cl=clusterName.get(find(vs), '') if vs in parent else '', hubn=(hubCount.get((regionOf.get(vs) or '') + '|' + hub, 1) if hub else 0)))

# 7. lộ trình chuyến
TN, tnIx, T = [], {}, {}
def ni(n):
    if n not in tnIx: tnIx[n] = len(TN); TN.append(n)
    return tnIx[n]
for code, t in trips.items():
    if not t['stops']: continue
    vIx = VEH.index(t['veh']) if t['veh'] in VEH else len(VEH) - 1
    st = [[ni(p[1]), p[2], round(p[3]), round(p[4]), p[5], p[6], p[7], round(p[8] or 0)] for p in sorted(t['path'], key=lambda p: p[0])]
    T[code] = [dayIx[t['d']], vIx, (1 if t['adhoc'] else 0) | (2 if t['inh'] else 0), t['km'], t['plate'], st]

# 8. toạ độ
GEO = {}
ssh = sheet(lwb, 'Station')
if ssh:
    sv = rows(ssh); hd = [str(x).strip() if x is not None else '' for x in sv[0]]
    cN, cLa, cLo = hd.index('Tên trạm'), hd.index('Vĩ độ'), hd.index('Kinh độ')
    pos, posL = {}, {}
    for r in sv[1:]:
        nm = str(r[cN] or '').strip()
        try: la = float(str(r[cLa]).replace(',', '.')); lo = float(str(r[cLo]).replace(',', '.'))
        except (TypeError, ValueError): continue
        if not nm or not math.isfinite(la) or not math.isfinite(lo) or not la or not lo: continue
        pos.setdefault(nm, [round(la * 1e5) / 1e5, round(lo * 1e5) / 1e5]); posL.setdefault(nm.lower(), pos[nm])
    want = set()
    for s in S:
        want.add(s['n'].strip())
        if s['h']: want.add(s['h'].strip())
        for n in s['soc']: want.add(n.strip())
    for n in TN: want.add(n.strip())
    for n in want:
        g = pos.get(n) or posL.get(n.lower())
        if g: GEO[n] = g

D = dict(v=8, win=92, warn=warn, pbt=pbt, dates=dates, dt=dtv, lh=[1 if d in lhSeen else 0 for d in dates], TY=VEH, S=S, T=T, TN=TN, GEO=GEO)
ref = re.search(r'^const REF=(.*);$', open(OLD, encoding='utf8').read(), re.M)[1]
open(OUT, 'w', encoding='utf8').write('const D=' + json.dumps(D, ensure_ascii=False, separators=(',', ':')) + ';\nconst REF=' + ref + ';\n')
lhd = [d for d in dates if d in lhSeen]
print(OUT, len(dates), 'ngày đơn', dates[0], '→', dates[-1], '· có chuyến', len(lhd), 'ngày', lhd[0], '→', lhd[-1], '·', len(S), 'điểm ·', len(T), 'chuyến ·', len(GEO), 'toạ độ')
